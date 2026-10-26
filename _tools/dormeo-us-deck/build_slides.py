#!/usr/bin/env python3
"""Dormeo US weekly business review -> small .pptx that Google Drive converts to Google Slides.
Usage: build_slides.py <deck_data.json> <extra.json> <notes.json> <out.pptx>
notes.json holds the week's written commentary (see RUNBOOK.md); every number in the tables and cards comes from the data files.
Charts are drawn with shapes (no embedded workbooks) and unused package parts are stripped so the
file stays small enough to upload through the Drive connector as base64."""
import json, sys, zipfile, re, io, datetime as dt
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

D = json.load(open(sys.argv[1])); X = json.load(open(sys.argv[2])); NOTES = json.load(open(sys.argv[3])); OUT = sys.argv[4]
NOTE_KEYS = ["week_story", "refunds_title", "refunds", "pacing_reco", "ads_context", "ads_strategy", "bid_moves",
             "coupons", "promotions", "events", "reviews_title", "reviews_body", "reviews_reco",
             "recommendations", "position", "next_steps", "data_ask"]
missing = [k_ for k_ in NOTE_KEYS if not NOTES.get(k_)]
if missing: sys.exit("notes.json is missing: " + ", ".join(missing))
N = NOTES
HERE = __file__.rsplit("/", 1)[0] + "/"
NAVY, BLUE, CREAM, LINE, MUTED = "0F172A", "2447FF", "F6F3EC", "E5E0D6", "6B6558"
GREEN, RED, AMBER, SKY, WHITE, INK2 = "16A34A", "DC2626", "D97706", "7EB8F7", "FFFFFF", "334155"
HF, BF = "Space Grotesk", "Inter"

def rgb(h): return RGBColor.from_string(h)
BRAND = D.get("brand", "Dormeo US")
CUR = D.get("money", "$")
def usd(v, d=0): return ("-" if v < 0 else "") + CUR + f"{abs(v):,.{d}f}"
def k(v): return (CUR + f"{v/1000:,.1f}K") if abs(v) >= 1000 else usd(v)
def pct(v, d=1): return f"{v:.{d}f}%"
def chg(c): return ("+" if c >= 0 else "") + f"{c*100:.0f}%"
def pts(c): return ("+" if c >= 0 else "") + f"{c:.1f} pts"

p = Presentation(); p.slide_width = Inches(13.333); p.slide_height = Inches(7.5)
blank = p.slide_layouts[6]
for l in list(p.slide_layouts):
    if l is not blank: p.slide_layouts.remove(l)

W = Week = D["week"]; WK = X["weeks"]; cur, prv = WK[-1], WK[-2]
END = dt.date.fromisoformat(W["end"]); TOTAL = 17

def box(s, x, y, w, h, fill=None, line=None, shape=MSO_SHAPE.RECTANGLE, lw=0.75):
    sh = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE: sh.adjustments[0] = 0.06
    if fill: sh.fill.solid(); sh.fill.fore_color.rgb = rgb(fill)
    else: sh.fill.background()
    if line: sh.line.color.rgb = rgb(line); sh.line.width = Pt(lw)
    else: sh.line.fill.background()
    sh.shadow.inherit = False
    return sh

def text(s, x, y, w, h, runs, size=12, color=NAVY, font=BF, bold=False, align="l", anchor="t", spacing=None):
    """runs: str, or list of paragraphs; a paragraph is str or list of (text, {opts})."""
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = tb.text_frame
    tf.word_wrap = True; tf.margin_left = tf.margin_right = Inches(0.04); tf.margin_top = tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
    paras = runs if isinstance(runs, list) else [runs]
    for i, para in enumerate(paras):
        pa = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        pa.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[align]
        if spacing: pa.space_after = Pt(spacing)
        for t, o in ([(para, {})] if isinstance(para, str) else para):
            r = pa.add_run(); r.text = t; f = r.font
            f.size = Pt(o.get("size", size)); f.bold = o.get("bold", bold); f.name = o.get("font", font)
            f.color.rgb = rgb(o.get("color", color))
    return tb

def header(s, kicker, title, n):
    box(s, 0, 0, 13.333, 7.5, fill=WHITE)
    box(s, 0.6, 0.55, 0.08, 0.62, fill=BLUE)
    text(s, 0.8, 0.42, 9.5, 0.3, kicker.upper(), size=10, color=BLUE, font=HF, bold=True)
    text(s, 0.8, 0.66, 10.5, 0.6, title, size=26, font=HF, bold=True)
    s.shapes.add_picture(HERE + "logo_s.png", Inches(11.15), Inches(0.5), Inches(1.6))
    box(s, 0.6, 7.0, 12.13, 0.012, fill=LINE)
    text(s, 0.6, 7.06, 9, 0.3, f"{BRAND} · Weekly Business Review · Week {W['num']} ({W['label']})", size=9, color=MUTED)
    text(s, 11.2, 7.06, 1.53, 0.3, f"{n} / {TOTAL}", size=9, color=MUTED, align="r")

def card(s, x, y, w, h, label, value, sub=None, subc=MUTED, fill=CREAM, vsize=26, vcolor=BLUE):
    box(s, x, y, w, h, fill=fill, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    text(s, x + 0.18, y + 0.14, w - 0.3, 0.3, label.upper(), size=9, color=MUTED, bold=True)
    text(s, x + 0.18, y + 0.42, w - 0.3, 0.6, value, size=vsize, color=vcolor, font=HF, bold=True)
    if sub: text(s, x + 0.18, y + h - 0.45, w - 0.3, 0.35, sub, size=10, color=subc)

def panel(s, x, y, w, h, title, body, fill=WHITE, line=LINE, tcolor=NAVY, size=11):
    box(s, x, y, w, h, fill=fill, line=line, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    text(s, x + 0.2, y + 0.15, w - 0.4, 0.3, title, size=12, font=HF, bold=True, color=tcolor)
    paras = body if isinstance(body, list) else [body]
    text(s, x + 0.2, y + 0.5, w - 0.4, h - 0.6, paras, size=size, color=INK2, spacing=5)

def table(s, x, y, w, rows, widths, size=10, rh=0.32, head_fill=NAVY, total_last=False, aligns=None, colors=None):
    nr, nc = len(rows), len(rows[0])
    t = s.shapes.add_table(nr, nc, Inches(x), Inches(y), Inches(w), Inches(rh * nr)).table
    for j, cw in enumerate(widths): t.columns[j].width = Inches(cw)
    for i, row in enumerate(rows):
        t.rows[i].height = Inches(rh)
        for j, v in enumerate(row):
            c = t.cell(i, j); c.margin_left = c.margin_right = Inches(0.06); c.margin_top = c.margin_bottom = Inches(0.03)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            c.fill.solid()
            head, tot = i == 0, total_last and i == nr - 1
            c.fill.fore_color.rgb = rgb(head_fill if head else (CREAM if tot else (WHITE if i % 2 else "FBFAF7")))
            tf = c.text_frame; tf.word_wrap = True; pa = tf.paragraphs[0]
            pa.alignment = PP_ALIGN.LEFT if (aligns and aligns[j] == "l") or (not aligns and j == 0) else PP_ALIGN.RIGHT
            r = pa.add_run(); r.text = str(v); f = r.font; f.size = Pt(size - (1 if head else 0)); f.name = BF
            f.bold = head or tot
            col = WHITE if head else NAVY
            if colors and not head and colors(i, j, v): col = colors(i, j, v)
            f.color.rgb = rgb(col)
    return t

def bars(s, x, y, w, h, items, color=BLUE, fmt=k, label_w=1.9, hi=None):
    """horizontal bars: items = [(label, value)]"""
    mx = max(v for _, v in items) or 1; n = len(items); rh = h / n
    for i, (lab, v) in enumerate(items):
        yy = y + i * rh
        text(s, x, yy, label_w, rh, lab, size=10, color=NAVY, anchor="m")
        bw = max(0.03, (w - label_w - 1.0) * v / mx)
        box(s, x + label_w, yy + rh * 0.22, bw, rh * 0.56, fill=(hi(i, v) if hi else color))
        text(s, x + label_w + bw + 0.06, yy, 1.0, rh, fmt(v), size=10, color=NAVY, bold=True, anchor="m")

def colbars(s, x, y, w, h, items, color=BLUE, fmt=k, highlight=None):
    """vertical bars: items = [(label, value)]"""
    mx = max(v for _, v in items) or 1; n = len(items); cw = w / n
    for i, (lab, v) in enumerate(items):
        bh = (h - 0.6) * v / mx; xx = x + i * cw
        box(s, xx + cw * 0.18, y + h - 0.3 - bh, cw * 0.64, bh, fill=(highlight if highlight and i == n - 1 else color))
        text(s, xx, y + h - 0.32 - bh - 0.28, cw, 0.26, fmt(v), size=9, bold=True, align="c")
        text(s, xx, y + h - 0.28, cw, 0.28, lab, size=9, color=MUTED, align="c")
    box(s, x, y + h - 0.3, w, 0.01, fill=LINE)

def slide(): return p.slides.add_slide(blank)
n = 0
def nxt():
    global n; n += 1; return n

# ---------------- numbers ----------------
rv, rvp = cur["revenue"], prv["revenue"]
asp, aspp = rv / cur["units"], rvp / prv["units"]
pac = D["pacing"]; mtd = X["mtd"]; ly = X["mtdLY"]
dim = pac["dim"]; days_left = dim - mtd["days"]
spend_pace = mtd["spend"] / mtd["days"] * dim
yoy = mtd["revenue"] / ly["revenue"] - 1
inv = D["inventory"]
strat = {a["name"]: a for a in D["adsByStrat"]}
br, ca = strat["Branded keywords"], strat["Category keywords"]
st = X["searchTerms"]; tg = X["targets"]; al = X["alerts"]

# 1 cover
s = slide(); nxt()
box(s, 0, 0, 13.333, 7.5, fill=NAVY)
box(s, 0, 0, 0.18, 7.5, fill=BLUE)
s.shapes.add_picture(HERE + "logo_dark_s.png", Inches(0.9), Inches(0.85), Inches(2.8))
text(s, 0.9, 2.35, 11, 0.4, "WEEKLY PERFORMANCE & GROWTH REVIEW", size=14, color=SKY, font=HF, bold=True)
text(s, 0.9, 2.8, 11, 1.2, BRAND, size=60, color=WHITE, font=HF, bold=True)
text(s, 0.9, 4.05, 11, 0.5, f"Amazon Business Review · Week {W['num']}", size=22, color=WHITE, font=HF)
text(s, 0.9, 4.7, 11, 0.4, f"Week ending {END.strftime('%A, %B')} {END.day}, {END.year}  ·  {W['label']}", size=14, color=SKY)
if D.get("currency", "USD") != "USD":   # e.g. Dormeo Canada: say which currency every amount is in
    text(s, 0.9, 5.2, 11, 0.4, f"All amounts in {D['currency']}", size=12, color=SKY)

# 2 agenda
s = slide(); header(s, "Today's agenda", "What we'll cover", nxt())
ag = [("01", "Sales overview", "Four-week channel performance, price and refunds"),
      ("02", "Sales pacing", f"{D['pacing']['monthLong']} month to date vs last month and last year"),
      ("03", "Amazon Ads results", "Spend, efficiency, brand vs category, keywords"),
      ("04", "Marketing", "Coupons, promotions and upcoming events"),
      ("05", "Inventory", "Stock position, weeks of supply, restock risk"),
      ("06", "Reviews & account health", "Ratings, alerts and Buy Box")]
for i, (num, t, d) in enumerate(ag):
    x = 0.6 + (i % 3) * 4.08; y = 1.75 + (i // 3) * 2.45
    box(s, x, y, 3.9, 2.2, fill=CREAM, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    text(s, x + 0.25, y + 0.2, 1.5, 0.6, num, size=30, color=BLUE, font=HF, bold=True)
    text(s, x + 0.25, y + 0.95, 3.4, 0.4, t, size=17, font=HF, bold=True)
    text(s, x + 0.25, y + 1.4, 3.4, 0.7, d, size=11, color=MUTED)

# 3 channel performance
s = slide(); header(s, "Sales overview", "Amazon channel performance — last 4 weeks", nxt())
rows = [["Week", "Ordered product sales", "Units ordered", "ASP", "Ad impressions", "Units refunded", "Refund rate"]]
for w in WK:
    rows.append([w["label"], usd(w["revenue"]), f"{w['units']:,}", usd(w["revenue"] / w["units"]), f"{w['impressions']:,}",
                 str(w["refundUnits"]), pct(w["refundRate"], 2)])
tot_r = sum(w["revenue"] for w in WK); tot_u = sum(w["units"] for w in WK); tot_ref = sum(w["refundUnits"] for w in WK)
rows.append(["4-week total", usd(tot_r), f"{tot_u:,}", usd(tot_r / tot_u), f"{sum(w['impressions'] for w in WK):,}", str(tot_ref), "—"])
table(s, 0.6, 1.55, 12.13, rows, [2.1, 2.2, 1.5, 1.4, 1.9, 1.5, 1.53], size=11, rh=0.4, total_last=True,
      colors=lambda i, j, v: (RED if j == 6 and i == 4 else None))
panel(s, 0.6, 4.25, 5.95, 2.55, "What happened this week",
      [f"Sales were {usd(rv)}, {chg(rv/rvp-1)} vs the prior week ({usd(rvp)}), on {cur['units']} units ({chg(cur['units']/prv['units']-1)})."] + N["week_story"])
panel(s, 6.78, 4.25, 5.95, 2.55, N["refunds_title"],
      [f"{cur['refundUnits']} units were refunded this week ({pct(cur['refundRate'],1)} refund rate), vs "
       f"{prv['refundUnits']} ({pct(prv['refundRate'],1)}) the week before."] + N["refunds"])

# 4 sales pacing
s = slide(); header(s, "Sales pacing", f"{pac['monthLong']} month to date", nxt())
MON = pac["month"]; MLONG = pac["monthLong"]
card(s, 0.6, 1.55, 2.9, 1.7, "MTD revenue", k(mtd["revenue"]), f"{MON} 1–{mtd['days']} · {mtd['orders']:,} orders")
card(s, 3.68, 1.55, 2.9, 1.7, "MTD ad spend", k(mtd["spend"]), f"ACoS {pct(mtd['acos'])}")
card(s, 6.76, 1.55, 2.9, 1.7, "MTD TACoS", pct(mtd["tacos"]), "ad spend ÷ total sales")
card(s, 9.84, 1.55, 2.89, 1.7, "vs last year", chg(yoy), f"{MON} 1–{mtd['days']}, {END.year-1}: {k(ly['revenue'])}", subc=GREEN if yoy >= 0 else RED, vcolor=GREEN if yoy >= 0 else RED)
colbars(s, 0.6, 3.5, 6.3, 3.3, [(f"{MON} {END.year-1}", pac["lastYear"]), (f"{pac['prevMonthName']} {END.year}", pac["prevMonth"]), (f"{MON} MTD", mtd["revenue"]),
                                 (f"{MON} pace", pac["forecast"])], color="C7D2FE", highlight=BLUE)
panel(s, 7.2, 3.5, 5.53, 1.45, "Full-month pace",
      [f"{MLONG} is pacing to about {k(pac['forecast'])}: {k(mtd['revenue'])} so far plus {days_left} days at "
       f"{k(pac['avg'])}/day" + (f" (the {MLONG} daily average without the {len(pac.get('eventDays', []))} event days)." if pac.get("eventDays") else f" (the {MLONG} daily average so far).")], size=10.5)
box(s, 7.2, 5.1, 5.53, 1.7, fill="FEF2F2", shape=MSO_SHAPE.ROUNDED_RECTANGLE)
text(s, 7.4, 5.22, 5.2, 0.3, "Recommendation", size=12, color=RED, font=HF, bold=True)
text(s, 7.4, 5.55, 5.2, 1.2, f"Pace is {chg(pac['forecast']/pac['prevMonth']-1)} vs {pac['prevMonthName']}. " + N["pacing_reco"], size=10.5, color=INK2)

# 5 pacing MTD ads
s = slide(); header(s, "Sales pacing", "Ad spend pacing — month to date", nxt())
card(s, 0.6, 1.55, 2.9, 1.7, "MTD ad spend", k(mtd["spend"]), f"{k(mtd['spend']/mtd['days'])} per day")
card(s, 3.68, 1.55, 2.9, 1.7, "Projected spend", k(spend_pace), f"if the daily rate holds to {MON} {dim}")
card(s, 6.76, 1.55, 2.9, 1.7, "MTD TACoS", pct(mtd["tacos"]), f"this week {pct(cur['tacos'])}")
card(s, 9.84, 1.55, 2.89, 1.7, "MTD ACoS", pct(mtd["acos"]), f"this week {pct(cur['acos'])}")
panel(s, 0.6, 3.5, 3.95, 3.3, "MTD performance",
      [f"{k(mtd['adSales'])} in ad sales on {k(mtd['spend'])} spend, a {mtd['adSales']/mtd['spend']:.1f}x return.",
       f"Ads drove {mtd['adSales']/mtd['revenue']*100:.0f}% of {MLONG} sales."])
panel(s, 4.69, 3.5, 3.95, 3.3, "Context", N["ads_context"])
panel(s, 8.78, 3.5, 3.95, 3.3, "This week's efficiency",
      [f"Spend {usd(cur['spend'])} ({chg(cur['spend']/prv['spend']-1)}), ad sales {usd(cur['adSales'])} ({chg(cur['adSales']/prv['adSales']-1)}).",
       f"ACoS {pct(cur['acos'])} ({pts(cur['acos']-prv['acos'])}); TACoS {pct(cur['tacos'])} ({pts(cur['tacos']-prv['tacos'])}).",
       f"CPC {usd(cur['cpc'],2)}, {cur['adOrders']} ad orders."])

# 6 brand vs category
s = slide(); header(s, "Amazon Ads", "Brand vs. category spend", nxt())
def r_(a): return [usd(a["spend"]), usd(a["sales"]), f"{a['sales']/a['spend']:.2f}x" if a["spend"] else "—", pct(a["acos"]*100) if a["sales"] else "—", str(a["orders"])]
rows = [["Strategy", "Ad spend", "Ad sales", "ROAS", "ACoS", "Orders"]]
for a in D["adsByStrat"]:
    if a["spend"] >= 50: rows.append([a["name"]] + r_(a))
cs = sum(a["spend"] for a in D["adsByStrat"]); cr = sum(a["sales"] for a in D["adsByStrat"]); co = sum(a["orders"] for a in D["adsByStrat"])
rows.append(["All campaigns", usd(cs), usd(cr), f"{cr/cs:.2f}x", pct(cs/cr*100), str(co)])
table(s, 0.6, 1.55, 7.4, rows, [2.9, 1.0, 1.0, 0.8, 0.85, 0.85], size=10.5, rh=0.38, total_last=True,
      colors=lambda i, j, v: (RED if j == 4 and v not in ("—",) and v.endswith("%") and float(v[:-1]) > 35 else None))
panel(s, 8.25, 1.55, 4.48, 2.7, "What it says",
      N["ads_strategy"], size=10.5)
panel(s, 8.25, 4.4, 4.48, 2.4, "Coverage note",
      ["Strategy is read from the campaign names (Brand / Category / Auto / Audience).",
       "Campaign totals come from the campaign report and can differ slightly from account totals."], size=10)

# 7 bid optimization
s = slide(); header(s, "Amazon Ads", "Bid optimization", nxt())
card(s, 0.6, 1.55, 2.9, 1.6, "Targets tracked", f"{tg['total']}", "keywords and product targets")
card(s, 3.68, 1.55, 2.9, 1.6, "Winners", f"{tg['winners']}", "converting at or under 35% ACoS", vcolor=GREEN)
card(s, 6.76, 1.55, 2.9, 1.6, "Need attention", f"{tg['needsAttention']}", f"{tg['highAcos']} high ACoS · {tg['noSales']} no sales", vcolor=RED)
card(s, 9.84, 1.55, 2.89, 1.6, "Low click rate", f"{tg['lowCtr']}", "seen but rarely clicked", vcolor=AMBER)
rows = [["Target", "Match", "Type", "Spend", "Sales", "ACoS", "Clicks", "Orders"]]
for t in tg["rows"]:
    rows.append([t[0], t[1], t[2], usd(t[3]), usd(t[4]), pct(t[5]) if t[5] is not None else "no sales", str(t[6]), str(t[7])])
table(s, 0.6, 3.35, 7.9, rows, [2.1, 0.8, 0.9, 0.9, 0.9, 0.9, 0.7, 0.7], size=10, rh=0.36,
      colors=lambda i, j, v: (RED if j == 5 else None))
panel(s, 8.75, 3.35, 3.98, 3.45, "Next moves",
      N["bid_moves"] + ["Impression share and bid-change history are not in the Sellerise report, so they aren't tracked here."], size=10.5)

# 8 brand keywords
s = slide(); header(s, "Amazon Ads", "Brand keywords — top customer searches", nxt())
rs = st["rows"]
for i, r in enumerate(rs[:6]):
    x = 0.6 + (i % 3) * 4.08; y = 1.5 + (i // 3) * 1.25
    box(s, x, y, 3.9, 1.1, fill=CREAM, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    text(s, x + 0.18, y + 0.1, 3.6, 0.3, f"“{r[0]}” · {r[1]}", size=10.5, bold=True)
    text(s, x + 0.18, y + 0.42, 3.6, 0.6, [[(usd(r[6]), {"size": 20, "font": HF, "bold": True, "color": BLUE}),
                                           (f"   {r[3]} clicks · {r[4]} orders · CVR {r[4]/r[3]*100:.1f}%", {"size": 10, "color": MUTED})]])
rows = [["Search term", "Ad", "Impressions", "Clicks", "Orders", "CVR", "Spend", "Sales", "ACoS"]]
for r in rs:
    rows.append([r[0], r[1], f"{r[2]:,}", str(r[3]), str(r[4]), pct(r[4]/r[3]*100), usd(r[5]), usd(r[6]), pct(r[5]/r[6]*100)])
table(s, 0.6, 4.1, 12.13, rows, [3.6, 0.6, 1.25, 0.9, 0.9, 0.9, 1.3, 1.3, 1.38], size=9.5, rh=0.3,
      colors=lambda i, j, v: (RED if j == 8 and float(v[:-1]) > 35 else None))

# 9 category keywords
s = slide(); header(s, "Amazon Ads", "Category keywords", nxt())
rows = [["Search term", "Ad", "Impressions", "Clicks", "Orders", "Spend", "Sales"]]
for r in st["category"]:
    rows.append([r[0], r[1], f"{r[2]:,}", str(r[3]), str(r[4]), usd(r[5]), usd(r[6])])
table(s, 0.6, 1.55, 12.13, rows, [4.73, 1.0, 1.6, 1.2, 1.2, 1.2, 1.2], size=11, rh=0.42)
text(s, 0.6, 1.7 + 0.42 * len(rows), 12.13, 0.5, f"Top category searches by ad sales, {W['label']} (Sponsored Products and Brands search-term report; only searches with at least one click are reported).", size=9, color=MUTED)

# 10 advertising performance daily
s = slide(); header(s, "Amazon Ads", "Advertising performance — daily, last 2 weeks", nxt())
rows = [["Date", "Day", "Total sales", "PPC spend", "PPC sales", "ROAS", "ACoS", "TACoS", "Orders", "Ad orders", "Organic sales", "Ad sales %"]]
def wkrow(lbl, ds):
    t = [sum(d[i] for d in ds) for i in range(1, 6)]
    return [lbl, "", usd(t[0]), usd(t[1]), usd(t[2]), f"{t[2]/t[1]:.2f}x", pct(t[1]/t[2]*100), pct(t[1]/t[0]*100), str(t[3]), str(t[4]), usd(t[0]-t[2]), pct(t[2]/t[0]*100, 0)]
for idx, dd in enumerate(X["daily"]):
    d0 = dt.date.fromisoformat(dd[0]); r, sp, sa, o, ao = dd[1:]
    rows.append([d0.strftime("%b %-d"), d0.strftime("%a"), usd(r), usd(sp), usd(sa), f"{sa/sp:.2f}x", pct(sp/sa*100), pct(sp/r*100),
                 str(o), str(ao), usd(max(0, r-sa)), pct(sa/r*100, 0)])
    if idx == 6: rows.append(wkrow(f"Week {prv['label']}", X["daily"][:7]))
rows.append(wkrow(f"Week {cur['label']}", X["daily"][7:]))
tb = table(s, 0.6, 1.45, 12.13, rows, [1.35, 0.6, 1.15, 1.05, 1.05, 0.8, 0.8, 0.8, 0.8, 0.9, 1.2, 0.83], size=8.5, rh=0.27,
           colors=lambda i, j, v: (RED if j == 6 and v.endswith("%") and float(v[:-1]) > 35 else None))
for ri in (8, 16):
    for j in range(12):
        c = tb.cell(ri, j); c.fill.solid(); c.fill.fore_color.rgb = rgb(CREAM)
        for r_ in c.text_frame.paragraphs[0].runs: r_.font.bold = True
text(s, 0.6, 6.5, 12.13, 0.45, f"Organic sales = total sales minus ad-attributed sales (floored at {CUR}0). Daily clicks and impressions aren't in the Sellerise daily feed; "
     f"weekly totals: {cur['clicks']:,} clicks / {cur['impressions']:,} impressions this week vs {prv['clicks']:,} / {prv['impressions']:,}.", size=9, color=MUTED)

# 11 marketing
s = slide(); header(s, "Marketing", "Coupons, promotions & upcoming events", nxt())
box(s, 0.6, 1.5, 12.13, 0.55, fill="EEF2FF", shape=MSO_SHAPE.ROUNDED_RECTANGLE)
text(s, 0.8, 1.5, 11.8, 0.55, "Data source: Sellerise promotion totals. Coupon and deal names aren't in the feed, so set-up details come from the account team.", size=10.5, color=BLUE, anchor="m")
panel(s, 0.6, 2.25, 3.95, 4.55, "Coupons", N["coupons"], size=10.5)
panel(s, 4.69, 2.25, 3.95, 4.55, "Promotions", N["promotions"], size=10.5)
panel(s, 8.78, 2.25, 3.95, 4.55, "Upcoming events", N["events"], size=10.5)

# 12 inventory
s = slide(); header(s, "Inventory", f"Stock position — week {inv['week']} inventory update ({inv['asOf']})", nxt())
st = D["invStatus"]  # from the inventory sheet only
card(s, 0.6, 1.5, 2.9, 1.55, "Out of stock", str(st["oos"]), f"of {st['selling']} selling SKUs, 0 at FBA", vcolor=RED)
card(s, 0.6, 3.2, 2.9, 1.55, "Low stock", str(st["low"]), "4 weeks of stock or less", vcolor=AMBER)
card(s, 0.6, 4.9, 2.9, 1.9, "Units inbound", f"{st['inbound']:,.0f}", "on the way to FBA")
text(s, 3.75, 1.5, 8.98, 0.3, "Weeks of supply at FBA (available ÷ weekly sales)", size=11, bold=True)
items = [(r["name"], min(r["weeksAvail"], 12)) for r in inv["rows"]]
raw = {r["name"]: r["weeksAvail"] for r in inv["rows"]}
bars(s, 3.75, 1.85, 8.98, 4.9, items, fmt=lambda v: (f"{v:.1f} wk" if v < 12 else "12+ wk"), label_w=2.0,
     hi=lambda i, v: RED if v < 1 else (AMBER if v < 3 else BLUE))

# 13 reviews
s = slide(); header(s, "Reviews & account health", "Rating & review trend", nxt())
box(s, 0.6, 1.55, 6.0, 5.25, fill=CREAM, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
text(s, 0.85, 1.8, 5.5, 0.4, "LAST 90 DAYS", size=11, color=AMBER, bold=True)
text(s, 0.85, 2.2, 5.5, 0.8, N["reviews_title"], size=24, font=HF, bold=True)
text(s, 0.85, 3.2, 5.5, 3.4, N["reviews_body"], size=11.5, color=INK2, spacing=8)
panel(s, 6.85, 1.55, 5.88, 5.25, "Recommendation", N["reviews_reco"], size=11)

# 14 alerts & buy box
s = slide(); header(s, "Reviews & account health", "Alerts & Buy Box — last 30 days", nxt())
card(s, 0.6, 1.5, 2.9, 1.55, "Total alerts", str(al["total"]), "open, incl. log events")
card(s, 0.6, 3.2, 2.9, 1.55, "Actionable", str(al["actionable"]), "warnings + critical", vcolor=AMBER)
card(s, 0.6, 4.9, 2.9, 1.9, "Critical", str(al["critical"]), "Buy Box, stock, inactive", vcolor=RED)
text(s, 3.75, 1.5, 8.98, 0.3, "Alerts by type", size=11, bold=True)
bt = sorted(al["byType"].items(), key=lambda kv: -kv[1])[:10]
bars(s, 3.75, 1.85, 8.98, 4.9, bt, fmt=lambda v: f"{v:,}", label_w=2.0,
     hi=lambda i, v: RED if bt[i][0] in ("Buy Box lost", "Out of stock", "Inactive") else BLUE)

# 15 recommendations
s = slide(); header(s, "Next steps", "This week's recommendations", nxt())
rows = [["Area", "Recommendation", "Owner", "Timing"]] + [list(r) for r in N["recommendations"][:8]]
table(s, 0.6, 1.5, 12.13, rows, [1.6, 7.73, 1.5, 1.3], size=11, rh=min(0.62, 5.3 / len(rows)), aligns=["l", "l", "l", "l"])

# 16 ad spend performance (cream)
s = slide(); box(s, 0, 0, 13.333, 7.5, fill=CREAM); n += 1
box(s, 0.6, 0.55, 0.08, 0.62, fill=BLUE)
text(s, 0.8, 0.42, 9, 0.3, "AMAZON ADS", size=10, color=BLUE, font=HF, bold=True)
text(s, 0.8, 0.66, 10.5, 0.6, "Ad spend performance — this week vs last", size=26, font=HF, bold=True)
s.shapes.add_picture(HERE + "logo_s.png", Inches(11.15), Inches(0.5), Inches(1.6))
for i, (lab, v, pv, f, better) in enumerate([("Ad spend", cur["spend"], prv["spend"], usd, "down"),
                                              ("Ad sales", cur["adSales"], prv["adSales"], usd, "up"),
                                              ("ACoS", cur["acos"], prv["acos"], pct, "down"),
                                              ("TACoS", cur["tacos"], prv["tacos"], pct, "down")]):
    pc = f in (pct,); dv = (v - pv) if pc else (v / pv - 1)
    good = (dv < 0) == (better == "down")
    card(s, 0.6 + i * 3.08, 1.55, 2.9, 1.7, lab, f(v), (pts(dv) if pc else chg(dv)) + f" vs {f(pv)}", subc=GREEN if good else RED, fill=WHITE)
x0 = 0.6
for i, (lab, a, b) in enumerate([("Ad spend", prv["spend"], cur["spend"]), ("Ad sales", prv["adSales"], cur["adSales"]),
                                 ("Total sales", prv["revenue"], cur["revenue"])]):
    xx = x0 + i * 4.1; mx = max(a, b); hh = 2.35
    text(s, xx, 3.45, 3.8, 0.3, lab, size=12, bold=True, font=HF)
    for jj, (v, col, nm) in enumerate([(a, "C7D2FE", prv["label"]), (b, BLUE, cur["label"])]):
        bh = hh * v / mx; bx = xx + 0.4 + jj * 1.5
        box(s, bx, 6.55 - bh, 1.1, bh, fill=col)
        text(s, bx - 0.2, 6.55 - bh - 0.3, 1.5, 0.28, k(v), size=10, bold=True, align="c")
        text(s, bx - 0.2, 6.6, 1.5, 0.28, nm, size=9, color=MUTED, align="c")
text(s, 0.6, 7.06, 9, 0.3, f"{BRAND} · Weekly Business Review · Week {W['num']} ({W['label']})", size=9, color=MUTED)
text(s, 11.2, 7.06, 1.53, 0.3, f"{n} / {TOTAL}", size=9, color=MUTED, align="r")

# 17 path forward
s = slide(); header(s, "Path forward", "Where we stand and what's next", nxt())
card(s, 0.6, 1.5, 2.9, 1.55, "MTD ad spend", k(mtd["spend"]))
card(s, 3.68, 1.5, 2.9, 1.55, "Projected spend", k(spend_pace))
card(s, 6.76, 1.5, 2.9, 1.55, "MTD ACoS", pct(mtd["acos"]))
card(s, 9.84, 1.5, 2.89, 1.55, f"Days left in {MON}", str(days_left))
panel(s, 0.6, 3.25, 3.95, 3.55, "Position", N["position"], size=10.5)
panel(s, 4.69, 3.25, 3.95, 3.55, "Next steps", [f"{i + 1}. {t}" for i, t in enumerate(N["next_steps"])], size=10.5)
panel(s, 8.78, 3.25, 3.95, 3.55, "Data ask", N["data_ask"], size=10.5)

assert n == TOTAL, n
buf = io.BytesIO(); p.save(buf)
# strip parts Google Slides doesn't need
zin = zipfile.ZipFile(io.BytesIO(buf.getvalue())); zout = zipfile.ZipFile(OUT, "w", zipfile.ZIP_DEFLATED, compresslevel=9)
DROP = {"docProps/thumbnail.jpeg", "ppt/printerSettings/printerSettings1.bin", "docProps/app.xml", "docProps/core.xml", "ppt/viewProps.xml", "ppt/presProps.xml"}
for it in zin.infolist():
    if it.filename in DROP: continue
    data = zin.read(it.filename)
    if it.filename == "[Content_Types].xml":
        data = re.sub(rb'<Default Extension="(jpeg|bin)"[^>]*/>', b"", data)
        data = re.sub(rb'<Override PartName="/(docProps/[^"]*|ppt/viewProps.xml|ppt/presProps.xml)"[^>]*/>', b"", data)
    if re.match(r"ppt/slides/slide\d+\.xml$", it.filename):
        data = re.sub(rb"<p:style>.*?</p:style>", b"", data)
        data = re.sub(rb'<p:txBody><a:bodyPr rtlCol="0" anchor="ctr"/><a:lstStyle/><a:p><a:pPr algn="ctr"/></a:p></p:txBody>', b"", data)
        data = re.sub(rb' name="[^"]*"', b' name=""', data)
        data = re.sub(rb' lang="[^"]*"| dirty="0"| rtlCol="0"', b"", data)
        data = data.replace(b'<a:effectLst/>', b'')
    if it.filename.endswith(".rels"):
        data = re.sub(rb'<Relationship [^>]*(thumbnail|printerSettings|docProps|viewProps|presProps)[^>]*/>', b"", data)
    zout.writestr(it, data)
zout.close()
import os; print("wrote", OUT, os.path.getsize(OUT), "bytes,", n, "slides")
