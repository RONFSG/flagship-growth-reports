#!/usr/bin/env python3
"""Keep ONE Google Slides presentation in step with the weekly deck, as real editable slides (same link every week).

The deck is first built with build_slides.py (.pptx). This script reads that file, describes every slide element
(shapes, text boxes, tables drawn as one text box per column, logos), compares it with what was last sent to the
presentation (gslides_state.json) and writes the Google Slides API requests that bring the presentation up to date:
  * a slide whose layout changed (or that doesn't exist yet) is deleted and rebuilt;
  * otherwise only what changed is sent: new text (with its styling), new size/position (e.g. bar lengths),
    new fill colours, new page background.

Usage:
  gslides_sync.py plan <deck.pptx> <out_dir> --existing <slide ids in the presentation, comma separated>
      [--force-geometry 3,7]  re-send position/size of every element on those slides (repair)
      writes out_dir/batch-NN.json (each a JSON list of requests; send each with update_presentation, in order)
      and out_dir/pending.json. Pass the ids from read_presentation (fields: slides.objectId).
  gslides_sync.py commit <out_dir> --done 1,2,...
      after the batches were sent, records the slides of the batches that succeeded into gslides_state.json.
"""
import hashlib, json, os, sys
from pptx import Presentation
from pptx.util import Emu
from pptx.enum.shapes import MSO_SHAPE_TYPE, MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

HERE = os.path.dirname(os.path.abspath(__file__))
STATE = os.path.join(HERE, "gslides_state.json")
LOGO_URLS = {"light_bg": "https://reports.flagshipgrowth.com/assets/fsg-logo-transparent.png",
             "dark_bg": "https://reports.flagshipgrowth.com/assets/fsg-logo-dark-trim.png"}
MAX_BATCH_CHARS = 45000
ALIGN = {PP_ALIGN.CENTER: "CENTER", PP_ALIGN.RIGHT: "END"}
# The Google Slides file is 720 x 405 pt (10 x 5.625 in); the deck is designed at 960 x 540 pt (13.333 x 7.5 in).
SCALE = 720 / 960
PT = lambda emu: round(int(emu) / 12700 * SCALE, 1)
J = lambda o: json.dumps(o, separators=(",", ":"), ensure_ascii=False)


def rgb(h):
    h = str(h); c = {}
    for k, v in (("red", h[0:2]), ("green", h[2:4]), ("blue", h[4:6])):
        f = round(int(v, 16) / 255, 3)
        if f:
            c[k] = f
    return c


def run_style(r):
    f = r.font; st = {}
    if f.name: st["f"] = f.name
    if f.size: st["s"] = round(f.size.pt, 1)
    if f.bold: st["b"] = 1
    try:
        if f.color.rgb is not None: st["c"] = str(f.color.rgb)
    except Exception:
        pass
    return st


def api_style(st):
    o, fl = {}, []
    if "f" in st: o["fontFamily"] = st["f"]; fl.append("fontFamily")
    if "s" in st: o["fontSize"] = {"magnitude": round(st["s"] * SCALE, 1), "unit": "PT"}; fl.append("fontSize")
    o["bold"] = bool(st.get("b")); fl.append("bold")
    if "c" in st: o["foregroundColor"] = {"opaqueColor": {"rgbColor": rgb(st["c"])}}; fl.append("foregroundColor")
    return o, ",".join(fl)


def frame_model(tf):
    """paragraphs -> [(align, [(text, style), ...]), ...]"""
    out = []
    for p in tf.paragraphs:
        runs = [(r.text, run_style(r)) for r in p.runs if r.text]
        out.append((ALIGN.get(p.alignment, "START"), runs))
    while out and not out[-1][1]:
        out.pop()
    return out


# ---------------------------------------------------------------- slide model
def model_slide(sl):
    els, bg = [], None
    k = 0
    for sh in sl.shapes:
        x, y, w, h = int(sh.left or 0), int(sh.top or 0), int(sh.width or 0), int(sh.height or 0)
        is_text = sh.has_text_frame and sh.text_frame.text.strip() != ""
        if x == 0 and y == 0 and w >= 12000000 and h >= 6800000 and not is_text:
            try:
                bg = str(sh.fill.fore_color.rgb) if sh.fill.type == 1 else None
            except Exception:
                bg = None
            continue
        k += 1
        if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
            els.append({"k": "img", "g": [x, y, w, h], "url": LOGO_URLS["dark_bg" if h and w / h > 7 else "light_bg"]})
            continue
        if getattr(sh, "has_table", False) and sh.has_table:
            els += model_table(sh)
            continue
        e = {"k": "text" if sh.shape_type == MSO_SHAPE_TYPE.TEXT_BOX else "shape", "g": [x, y, w, h]}
        if e["k"] == "shape":
            try:
                e["t"] = "ROUND_RECTANGLE" if sh.auto_shape_type == MSO_SHAPE.ROUNDED_RECTANGLE else "RECTANGLE"
            except Exception:
                e["t"] = "RECTANGLE"
            try:
                e["fill"] = str(sh.fill.fore_color.rgb) if sh.fill.type == 1 else None
            except Exception:
                e["fill"] = None
            try:
                e["line"] = str(sh.line.color.rgb) if sh.line.fill.type == 1 else None
            except Exception:
                e["line"] = None
        if sh.has_text_frame:
            e["va"] = {MSO_ANCHOR.MIDDLE: "MIDDLE", MSO_ANCHOR.BOTTOM: "BOTTOM"}.get(sh.text_frame.vertical_anchor, "TOP")
            e["p"] = frame_model(sh.text_frame)
            sa = sh.text_frame.paragraphs[0].space_after
            if sa and len(e["p"]) > 1:
                e["sa"] = round(sa.pt, 1)
        els.append(e)
    return {"bg": bg, "els": els}


def model_table(sh):
    """A table becomes: a filled band per non-white row, then one text box per column (one line per row)."""
    t = sh.table; nr, nc = len(t.rows), len(t.columns)
    x0, y0 = int(sh.left), int(sh.top)
    rh = int(t.rows[0].height); W = sum(int(c.width) for c in t.columns)
    out = []
    for i in range(nr):
        try:
            fc = str(t.cell(i, 0).fill.fore_color.rgb)
        except Exception:
            fc = None
        if fc and fc.upper() != "FFFFFF":
            out.append({"k": "shape", "t": "RECTANGLE", "g": [x0, y0 + i * rh, W, rh], "fill": fc, "line": None})
    # one font size for every line of the table, so all columns share one line height and rows line up
    sizes = [r.font.size.pt for i in range(nr) for j in range(nc) for p in t.cell(i, j).text_frame.paragraphs for r in p.runs if r.font.size]
    tsize = round(max(sizes), 1) if sizes else 10
    cells = []  # cells[j][i] = (align, text, style)
    for j in range(nc):
        colcells = []
        for i in range(nr):
            fm = frame_model(t.cell(i, j).text_frame)
            al = fm[0][0] if fm else "START"
            runs = [r for p in fm for r in p[1]]
            st = dict(runs[0][1]) if runs else {}
            st["s"] = tsize
            colcells.append((al, " ".join(r[0] for r in runs) or " ", st))
        cells.append(colcells)
    # Slides text boxes keep a 0.1in inset on each side; column boxes get that back as extra width so short
    # headers like "Orders" don't wrap. If any cell would still wrap, lines in a column box would no longer
    # sit on their rows, so that table is drawn one box per cell instead (vertically centred in its row).
    pad = int(INSET_PT * 2 / SCALE * 12700)
    fits = all(est_width(txt, tsize * SCALE) <= PT(int(t.columns[j].width) + pad) - INSET_PT * 2
               for j in range(nc) for _, txt, _ in cells[j])
    x = x0
    for j, col in enumerate(t.columns):
        cw = int(col.width)
        if fits:
            data_al = cells[j][1][0] if nr > 1 else cells[j][0][0]
            gx, gw = (x - pad, cw + pad) if data_al == "END" else ((x - pad // 2, cw + pad) if data_al == "CENTER" else (x, cw + pad))
            out.append({"k": "text", "g": [gx, y0, gw, rh * nr], "va": "TOP",
                        "p": [(al, [(txt, st)]) for al, txt, st in cells[j]], "rowh": rh, "fs": tsize})
        else:
            for i, (al, txt, st) in enumerate(cells[j]):
                out.append({"k": "text", "g": [x, y0 + i * rh, cw, rh], "va": "MIDDLE", "p": [(al, [(txt, st)])]})
        x += cw
    return out


INSET_PT = 7.2  # Google Slides' default text-box inset on each side


def est_width(text, size_pt):
    """Rough rendered width in points (Inter/Space Grotesk average ~0.55em per character)."""
    return len(text) * size_pt * 0.55


def layout_sig(m):
    """What must match for an in-place update: element kinds/types/images in order."""
    return hashlib.md5(J([[e["k"], e.get("t"), e.get("url"), e["g"] if e["k"] == "img" else None] for e in m["els"]]).encode()).hexdigest()[:12]


# ---------------------------------------------------------------- request builders
def el_props(page, g):
    x, y, w, h = g
    return {"pageObjectId": page, "size": {"width": {"magnitude": max(PT(w), 0.1), "unit": "PT"}, "height": {"magnitude": max(PT(h), 0.1), "unit": "PT"}},
            "transform": {"scaleX": 1, "scaleY": 1, "translateX": PT(x), "translateY": PT(y), "unit": "PT"}}


def text_reqs(oid, e, had_text):
    paras = e.get("p") or []
    text = "\n".join("".join(r[0] for r in runs) for _, runs in paras)
    reqs = []
    if had_text:
        reqs.append({"deleteText": {"objectId": oid, "textRange": {"type": "ALL"}}})
    if not text.strip():
        return reqs
    reqs.append({"insertText": {"objectId": oid, "text": text, "insertionIndex": 0}})
    # styles: the most common run style for ALL, then ranges that differ
    spans, pos = [], 0
    for pi, (al, runs) in enumerate(paras):
        for t, st in runs:
            spans.append((pos, pos + len(t), st)); pos += len(t)
        pos += 1
    counts = {}
    for s, e2, st in spans:
        counts[J(st)] = counts.get(J(st), 0) + (e2 - s)
    base = json.loads(max(counts, key=counts.get)) if counts else {}
    o, fl = api_style(base)
    reqs.append({"updateTextStyle": {"objectId": oid, "style": o, "fields": fl, "textRange": {"type": "ALL"}}})
    merged = []
    for s, e2, st in spans:
        if st == base or e2 <= s:
            continue
        if merged and merged[-1][2] == st and merged[-1][1] >= s - 1:
            merged[-1][1] = e2
        else:
            merged.append([s, e2, st])
    for s, e2, st in merged:
        o, fl = api_style(st)
        reqs.append({"updateTextStyle": {"objectId": oid, "style": o, "fields": fl, "textRange": {"type": "FIXED_RANGE", "startIndex": s, "endIndex": min(e2, len(text))}}})
    # paragraph alignment (+ row spacing for table columns)
    aligns = [al for al, _ in paras]
    ps, pf = {}, []
    common = max(set(aligns), key=aligns.count)
    if common != "START" or e.get("rowh"):
        ps["alignment"] = common; pf.append("alignment")
    if e.get("sa"):
        ps["spaceBelow"] = {"magnitude": round(e["sa"] * SCALE, 1), "unit": "PT"}; pf.append("spaceBelow")
    if e.get("rowh"):
        fs = e.get("fs") or base.get("s", 10)
        ps.update({"lineSpacing": round(PT(e["rowh"]) / (fs * SCALE * 1.2) * 100, 1), "spaceAbove": {"magnitude": 0, "unit": "PT"}, "spaceBelow": {"magnitude": 0, "unit": "PT"}})
        pf += ["lineSpacing", "spaceAbove", "spaceBelow"]
    if pf:
        reqs.append({"updateParagraphStyle": {"objectId": oid, "style": ps, "fields": ",".join(pf), "textRange": {"type": "ALL"}}})
    pos = 0
    for al, runs in paras:
        ln = sum(len(r[0]) for r in runs)
        if al != common and ln:
            reqs.append({"updateParagraphStyle": {"objectId": oid, "style": {"alignment": al}, "fields": "alignment", "textRange": {"type": "FIXED_RANGE", "startIndex": pos, "endIndex": pos + ln}}})
        pos += ln + 1
    return reqs


def shape_props(oid, e, va_needed):
    props, fl = {}, []
    if e["k"] == "shape":
        props["shapeBackgroundFill"] = {"solidFill": {"color": {"rgbColor": rgb(e["fill"])}}} if e.get("fill") else {"propertyState": "NOT_RENDERED"}
        fl.append("shapeBackgroundFill")
        props["outline"] = {"outlineFill": {"solidFill": {"color": {"rgbColor": rgb(e["line"])}}}, "weight": {"magnitude": 0.75, "unit": "PT"}} if e.get("line") else {"propertyState": "NOT_RENDERED"}
        fl.append("outline")
    if va_needed:
        props["contentAlignment"] = e.get("va", "TOP"); fl.append("contentAlignment")
    return [{"updateShapeProperties": {"objectId": oid, "shapeProperties": props, "fields": ",".join(fl)}}] if fl else []


def build_slide(page, idx, m):
    reqs = [{"createSlide": {"objectId": page, "insertionIndex": idx, "slideLayoutReference": {"predefinedLayout": "BLANK"}}}]
    if m["bg"] and m["bg"].upper() != "FFFFFF":
        reqs.append({"updatePageProperties": {"objectId": page, "pageProperties": {"pageBackgroundFill": {"solidFill": {"color": {"rgbColor": rgb(m["bg"])}}}}, "fields": "pageBackgroundFill.solidFill.color"}})
    for i, e in enumerate(m["els"], 1):
        oid = f"{page}e{i}"
        if e["k"] == "img":
            reqs.append({"createImage": {"objectId": oid, "url": e["url"], "elementProperties": el_props(page, e["g"])}})
            continue
        reqs.append({"createShape": {"objectId": oid, "shapeType": "TEXT_BOX" if e["k"] == "text" else e["t"], "elementProperties": el_props(page, e["g"])}})
        has_text = bool(e.get("p")) and any(r[0].strip() for _, runs in e["p"] for r in runs)
        reqs += shape_props(oid, e, has_text and (e["k"] == "shape" or e.get("va", "TOP") != "TOP"))
        reqs += text_reqs(oid, e, False)
    return reqs


BASE_EMU = 3000000


def update_slide(page, old, m, force_geo=False):
    reqs = []
    if (old["bg"] or "FFFFFF") != (m["bg"] or "FFFFFF"):
        reqs.append({"updatePageProperties": {"objectId": page, "pageProperties": {"pageBackgroundFill": {"solidFill": {"color": {"rgbColor": rgb(m["bg"] or "FFFFFF")}}}}, "fields": "pageBackgroundFill.solidFill.color"}})
    for i, (o, e) in enumerate(zip(old["els"], m["els"]), 1):
        oid = f"{page}e{i}"
        if e["k"] != "img" and (e["g"] != o["g"] or force_geo):
            # Slides stores every shape/text box at a 3,000,000 EMU base size and keeps the real size in the
            # transform's scale, so an ABSOLUTE transform's scale = wanted size / 3,000,000 EMU. (Images have a
            # pixel-based base size; their geometry is part of the layout signature, so a moved image rebuilds the slide.)
            sx = max(PT(e["g"][2]), 0.1) * 12700 / BASE_EMU
            sy = max(PT(e["g"][3]), 0.1) * 12700 / BASE_EMU
            reqs.append({"updatePageElementTransform": {"objectId": oid, "applyMode": "ABSOLUTE",
                         "transform": {"scaleX": round(sx, 5), "scaleY": round(sy, 5), "translateX": PT(e["g"][0]), "translateY": PT(e["g"][1]), "unit": "PT"}}})
        if e["k"] == "shape" and (e.get("fill") != o.get("fill") or e.get("line") != o.get("line")):
            reqs += shape_props(oid, e, False)
        if e["k"] != "img" and (J(e.get("p")) != J(o.get("p")) or e.get("sa") != o.get("sa")):
            had = bool(o.get("p")) and any(r[0].strip() for _, runs in o["p"] for r in runs)
            reqs += text_reqs(oid, e, had)
    return reqs


# ---------------------------------------------------------------- main
def plan(pptx, out, existing, force=()):
    os.makedirs(out, exist_ok=True)
    state = json.load(open(STATE)) if os.path.exists(STATE) else {"slides": []}
    models = [model_slide(s) for s in Presentation(pptx).slides]
    old = state.get("slides", [])
    known = {s["page"] for s in old}
    per_slide, new_slides, extra_delete = [], [], [x for x in existing if x not in known]
    for i, m in enumerate(models):
        o = old[i] if i < len(old) else None
        if o and o["page"] in existing and layout_sig(o) == layout_sig(m):
            page = o["page"]
            for e, oe in zip(m["els"], o["els"]):
                e["created"] = oe.get("created", oe["g"])
            reqs = update_slide(page, o, m, force_geo=(i + 1) in force)
        else:
            gen = (o["gen"] + 1) if o else 1
            page = f"s{i + 1:02d}g{gen}"
            if o and o["page"] in existing:
                extra_delete.append(o["page"])
            for e in m["els"]:
                e["created"] = e["g"]
            reqs = build_slide(page, i, m)
        per_slide.append(reqs)
        new_slides.append({"page": page, "gen": int(page.split("g")[-1]), "sig": layout_sig(m), "bg": m["bg"], "els": m["els"]})
    # slides beyond this week's count are deleted too
    for o in old[len(models):]:
        if o["page"] in existing:
            extra_delete.append(o["page"])
    # deletes go right after the first new/updated slide, so the presentation never has zero slides
    head = [{"deleteObject": {"objectId": x}} for x in extra_delete]
    per_slide[0] = per_slide[0] + head
    batches, cur, owners, cur_own = [], [], [], []
    for i, reqs in enumerate(per_slide):
        if cur and len(J(cur)) + len(J(reqs)) > MAX_BATCH_CHARS:
            batches.append(cur); owners.append(cur_own); cur, cur_own = [], []
        cur += reqs; cur_own.append(i)
    if cur or not batches:
        batches.append(cur); owners.append(cur_own)
    order = [{"updateSlidesPosition": {"slideObjectIds": [s["page"] for s in new_slides], "insertionIndex": 0}}]
    batches[-1] += order
    for f in os.listdir(out):
        if f.startswith("batch-"):
            os.remove(os.path.join(out, f))
    total = 0
    for n, b in enumerate(batches, 1):
        s = J(b); total += len(s)
        open(os.path.join(out, f"batch-{n:02d}.json"), "w").write(s)
    json.dump({"slides": new_slides, "owners": owners, "deleted": extra_delete}, open(os.path.join(out, "pending.json"), "w"))
    changed = sum(1 for r in per_slide if r)
    print(f"{len(models)} slides ({changed} with changes), {sum(len(b) for b in batches)} requests in {len(batches)} batches, {total:,} chars")
    for n, b in enumerate(batches, 1):
        print(f"  batch-{n:02d}.json: slides {', '.join(str(i + 1) for i in owners[n - 1])} · {len(b)} requests · {len(J(b)):,} chars")


def commit(out, done):
    p = json.load(open(os.path.join(out, "pending.json")))
    state = json.load(open(STATE)) if os.path.exists(STATE) else {"slides": []}
    old = state.get("slides", [])
    ok = set()
    for n in done:
        ok.update(p["owners"][n - 1])
    slides = []
    for i, s in enumerate(p["slides"]):
        if i in ok:
            slides.append(s)
        elif i < len(old):
            slides.append(old[i])
    state["slides"] = slides
    json.dump(state, open(STATE, "w"), separators=(",", ":"))
    print(f"recorded {len(ok)} of {len(p['slides'])} slides in gslides_state.json")


if __name__ == "__main__":
    cmd = sys.argv[1]
    if cmd == "plan":
        ex = sys.argv[sys.argv.index("--existing") + 1].split(",") if "--existing" in sys.argv else []
        fg = sys.argv[sys.argv.index("--force-geometry") + 1].split(",") if "--force-geometry" in sys.argv else []
        plan(sys.argv[2], sys.argv[3], [x for x in ex if x], {int(x) for x in fg if x})
    elif cmd == "commit":
        commit(sys.argv[2], [int(x) for x in sys.argv[sys.argv.index("--done") + 1].split(",") if x])
