#!/usr/bin/env python3
"""One command per step of the weekly deck, for any brand in _tools/deck-schedule.json.

Every brand uses the same 17-slide design (build_slides.py); each week only the data and commentary change.
This driver replaces the DECK_BRAND juggling and the long hand-run step lists, so a run needs fewer agent steps
and fewer Sellerise calls. Run from anywhere; <slug> is the brand's slug in deck-schedule.json (e.g. dormeo-us).

  deck.py brands                         enabled brands (monday + sunday lists) and whether each has a runbook
  deck.py index                          writes decks.json at the site root: the brands the Apps Script builds editable
                                         Google Slides for (enabled, with a runbook and a published deck). Run it when a
                                         brand is enabled or disabled, then push decks.json.
  deck.py tuesday <slug> [--dry-run]     the whole Tuesday inventory update up to Google Slides: pull, refresh the
                                         sheet data, skip if the sheet isn't newer than the published slide 12,
                                         rebuild with the week's saved inputs, publish, commit, push.
                                         Prints one JSON line: status, sheetDate, oos, low, image to check.
  deck.py extra-plan <slug>              Monday: the Sellerise calls still needed for extra.json (after deck_week.py
                                         build). The 4-week table, weekly clicks/impressions and last week's daily rows
                                         come from files already saved, so they are not pulled again.
  deck.py extra-build <slug> <pulled.json>   Monday: merge those pulls with the saved data into /tmp/deck/<slug>/extra.json
  deck.py build <slug>                   Monday: compute + build + render, and one contact-sheet image of all 17
                                         slides to look at (instead of 17 separate images)
  deck.py publish <slug>                 Monday: publish, save the week's extra.json/notes.json, commit, push
  deck.py slides <slug> <slide ids> [--only 1,12]
                                         LEGACY picture copy, not used any more: the brand's Google Slides file is rebuilt
                                         as editable slides by the Apps Script; never send pictures into it.
                                         the Google Slides requests (pass unchanged to update_presentation).
                                         --only replaces just those slides (Tuesday: 1,12), a much smaller call.

Work files go to /tmp/deck/<slug>/ (notes.json is read from there too).
"""
import datetime as dt, glob, json, os, re, shutil, subprocess, sys

HERE = os.path.dirname(os.path.abspath(__file__))
TOOLS = os.path.dirname(HERE)
REPO = os.path.dirname(TOOLS)
DASH = os.path.join(TOOLS, "dormeo-us-dashboard")
SCHED = json.load(open(os.path.join(TOOLS, "deck-schedule.json")))
D = dt.date.fromisoformat
MS = "Jan Feb Mar Apr May Jun Jul Aug Sep Oct Nov Dec".split()


def die(msg, **kw):
    print(json.dumps({"status": "stopped", "reason": msg, **kw}))
    sys.exit(2)


def entry(slug):
    for e in SCHED.get("monday", []) + SCHED.get("sunday", []):
        if e["slug"] == slug: return e
    die(f"{slug} is not in deck-schedule.json")


class Brand:
    def __init__(self, slug):
        self.slug, self.e = slug, entry(slug)
        self.deck = os.path.join(TOOLS, f"{slug}-deck")              # the brand's settings, data and saved weeks
        if not os.path.isdir(self.deck): die(f"{slug} has no _tools/{slug}-deck folder yet")
        self.us = slug == "dormeo-us"                                 # the original brand: its data lives with its dashboard
        self.data = DASH if self.us else self.deck
        self.config = os.path.join(self.deck, "deck_config.json")
        self.site = os.path.join(REPO, slug)
        self.tmp = os.path.join("/tmp/deck", slug); os.makedirs(self.tmp, exist_ok=True)
        self.env = {k: v for k, v in os.environ.items() if k != "DECK_BRAND"}
        if not self.us: self.env["DECK_BRAND"] = self.deck

    def run(self, args, cwd=HERE, ok=True):
        r = subprocess.run(args, cwd=cwd, env=self.env, capture_output=True, text=True, timeout=600)
        if ok and r.returncode: die(f"{' '.join(map(str, args[:2]))} failed", out=(r.stdout + r.stderr)[-800:])
        return r.stdout

    def manifest(self):
        p = os.path.join(self.site, "deck", "manifest.json")
        return json.load(open(p)) if os.path.exists(p) else None

    def weeks(self):
        return sorted(d for d in glob.glob(os.path.join(self.deck, "weeks", "*"))
                      if os.path.exists(os.path.join(d, "extra.json")) and os.path.exists(os.path.join(d, "notes.json")))

    # paths this brand's run may commit; nothing else ever gets staged
    def commit_paths(self):
        p = [os.path.join(self.site, "deck"), self.deck]
        if self.us: p.append(os.path.join(DASH, "data", "inv_sheet.json"))
        return p


def git(*a, ok=True):
    r = subprocess.run(["git", *a], cwd=REPO, capture_output=True, text=True, timeout=300)
    if ok and r.returncode: die(f"git {a[0]} failed", out=(r.stdout + r.stderr)[-800:])
    return r.stdout


def commit_push(b, msg):
    git("add", "--", *b.commit_paths())
    if not git("diff", "--cached", "--name-only").strip(): return "nothing to commit"
    git("commit", "-q", "-m", msg + "\n\nCo-Authored-By: Claude <noreply@anthropic.com>")
    if subprocess.run(["git", "push", "-q", "origin", "main"], cwd=REPO, capture_output=True).returncode:
        git("pull", "-q", "--rebase", "--autostash", "origin", "main"); git("push", "-q", "origin", "main")
    return git("log", "--oneline", "-1").strip()


def restore_unstaged(paths):
    """Put back files the sheet refresh touched but this run must not commit (e.g. Dormeo US's dashboard rec.json)."""
    for p in paths:
        changed = git("diff", "--name-only", "--", p).split()
        if changed: git("checkout", "--", *changed)


def sheet_date(s):  # "10/26/2026" -> date
    m, d, y = map(int, s.split("/")); return dt.date(y, m, d)


def published_inventory_date(b, man):
    cur = man["current"]
    if cur.get("inventoryAsOf"): return cur["inventoryAsOf"]
    from pptx import Presentation                                    # decks published before inventoryAsOf existed
    for s in Presentation(os.path.join(b.site, "deck", cur["pptx"])).slides:
        for sh in s.shapes:
            if sh.has_text_frame:
                m = re.search(r"inventory update \((\d+/\d+/\d{4})\)", sh.text_frame.text)
                if m: return m.group(1)
    return None


def build_deck(b, extra, notes, pages=None):
    dd = os.path.join(b.tmp, "deck_data.json"); px = os.path.join(b.tmp, "deck.pptx")
    b.run([sys.executable, "compute.py", b.data, dd, b.config])
    b.run([sys.executable, "build_slides.py", dd, extra, notes, px])
    b.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", b.tmp, px], cwd=b.tmp)
    for f in glob.glob(os.path.join(b.tmp, "s-*.jpg")): os.remove(f)
    rng = ["-f", str(pages[0]), "-l", str(pages[-1])] if pages else []
    b.run(["pdftoppm", *rng, "-r", "60", "-jpeg", os.path.join(b.tmp, "deck.pdf"), os.path.join(b.tmp, "s")], cwd=b.tmp)
    return dd, px


# ---------------------------------------------------------------- commands
def cmd_index():
    out = []
    for day in ("monday", "sunday"):
        for e in SCHED.get(day, []):
            if not e.get("enabled") or e.get("terminated") or not os.path.exists(os.path.join(REPO, e["runbook"])): continue
            if not os.path.exists(os.path.join(REPO, e["slug"], "deck", "manifest.json")): continue
            out.append({"name": e["brand"], "slug": e["slug"], "manifest": f"https://reports.flagshipgrowth.com/{e['slug']}/deck/manifest.json"})
    json.dump({"_about": "Brands whose weekly deck the Apps Script (_tools/apps-script) turns into editable Google Slides. "
                         "Written by _tools/dormeo-us-deck/deck.py index from deck-schedule.json.", "brands": out},
              open(os.path.join(REPO, "decks.json"), "w"), indent=1)
    print(json.dumps({"status": "ok", "brands": [b["slug"] for b in out]}))


def cmd_brands():
    for day in ("monday", "sunday"):
        for e in SCHED.get(day, []):
            if e.get("enabled"):
                print(day, e["slug"], "runbook" if os.path.exists(os.path.join(REPO, e["runbook"])) else "NO RUNBOOK")


def cmd_tuesday(slug, dry):
    b = Brand(slug)
    if not dry: git("pull", "-q", "--rebase", "--autostash", "origin", "main")
    wk = b.weeks()
    if not wk: die("no saved week: the Monday deck never published", brand=b.e["brand"])
    end = os.path.basename(wk[-1]); man = b.manifest()
    if not man or man["current"]["end"] != end:
        die(f"Monday deck for week ending {end} is not the published one", brand=b.e["brand"])
    out = b.run([sys.executable, "parse_sheet_pages.py", REPO, os.path.join(b.data, "data")], cwd=DASH)
    m = re.search(r"inventory (\d+/\d+/\d{4}) wk (\d+)", out)
    if not m: die("could not read the inventory sheet date", out=out[-400:])
    new, pub = m.group(1), published_inventory_date(b, man)
    if pub and sheet_date(new) <= sheet_date(pub):
        restore_unstaged([os.path.join(b.data, "data")])
        print(json.dumps({"status": "skipped", "brand": b.e["brand"], "reason": f"inventory sheet not updated yet (still {new})"}))
        return
    dd, _ = build_deck(b, os.path.join(wk[-1], "extra.json"), os.path.join(wk[-1], "notes.json"), pages=[12])
    st = json.load(open(dd))["invStatus"]
    res = {"status": "dry-run" if dry else "published", "brand": b.e["brand"], "weekEnd": end, "sheetDate": new,
           "sheetWeek": int(m.group(2)), "oos": st["oos"], "low": st["low"], "selling": st["selling"], "inbound": st["inbound"],
           "check": glob.glob(os.path.join(b.tmp, "s-*.jpg"))[0], "googleSlidesId": b.e.get("googleSlidesId")}
    if not dry:
        b.run([sys.executable, "publish_deck.py", os.path.join(b.tmp, "deck.pptx"), dd, b.site])
        res["commit"] = commit_push(b, f"{b.e['brand']} weekly deck: week {man['current']['week']} inventory update")
        restore_unstaged([os.path.join(b.data, "data")])
        res["next"] = "nothing: the Apps Script rebuilds the editable Google Slides file within ~30 minutes"
    else:
        restore_unstaged([os.path.join(b.data, "data")])
    print(json.dumps(res))


def mondays(b):
    w = json.load(open(os.path.join(b.deck, "wk", "deckweek.json")))
    mon = D(w["weekStart"]); return [mon - dt.timedelta(days=7 * i) for i in (3, 2, 1, 0)]


def label(mon):
    sun = mon + dt.timedelta(days=6)
    return f"{MS[mon.month-1]} {mon.day} – {sun.day if sun.month == mon.month else MS[sun.month-1] + ' ' + str(sun.day)}"


def cmd_extra_plan(slug):
    b = Brand(slug); M = mondays(b); mon = M[-1]; sun = mon + dt.timedelta(days=6)
    wkdir = os.path.join(b.deck, "wk")
    bj = json.load(open(os.path.join(b.deck, "brand.json"))) if os.path.exists(os.path.join(b.deck, "brand.json")) else {}
    base = {"accountSearch": bj.get("account", "Dormeo"), "marketplaces": bj.get("marketplace", "US")}
    am = {**base, "targetCurrency": bj["currency"]} if bj.get("currency") else base
    calls, missing = [], []
    for m in M:
        if not os.path.exists(os.path.join(wkdir, f"meta_{m}.json")): missing.append(str(m))
    prior = os.path.join(b.deck, "weeks", str(mon - dt.timedelta(days=1)), "extra.json")
    have_prior = os.path.exists(prior) and len(json.load(open(prior)).get("daily", [])) >= 14
    first = mon if have_prior else mon - dt.timedelta(days=7)
    y = sun; m1 = y.replace(day=1); ly = lambda d: d.replace(year=d.year - 1)
    calls += [
        {"key": "mtd", "tool": "get_account_metrics", "args": {**am, "groupBy": "none", "start": str(m1), "end": str(y)},
         "keep": "revenue, units, orders, ppc.spend->spend, ppc.sales->adSales, ppc.acos->acos, ppc.tacos->tacos, refundRate; add through (end date) and days (end day of month)"},
        {"key": "mtdLY", "tool": "get_account_metrics", "args": {**am, "groupBy": "none", "start": str(ly(m1)), "end": str(ly(y))}, "keep": "same as mtd"},
        {"key": "daily", "tool": "get_account_metrics", "args": {**am, "groupBy": "day", "start": str(first), "end": str(sun)},
         "keep": "[date, revenue, ppcSpend, ppcSales, orders, ppcOrders] per day" + ("" if have_prior else " (14 days: last week's saved rows are missing)")},
        {"key": "searchTerms", "tool": "get_ads_search_term_insights", "args": {**base, "start": str(mon), "end": str(sun), "mode": "top_sales", "limit": 25},
         "keep": "same shape as last week's extra.json searchTerms"},
        {"key": "targets", "tool": "get_ads_targeting_insights", "args": {**base, "start": str(mon), "end": str(sun), "mode": "needs_attention", "limit": 10},
         "keep": "same shape as last week's extra.json targets"},
        {"key": "reviews", "tool": "get_review_insights", "args": {**base, "start": "last 90 days", "end": "last 90 days"}, "keep": "{available:false} if no rows"},
        {"key": "alerts", "tool": "get_alert_summary", "args": {**base, "start": "last 30 days", "end": "last 30 days"}, "keep": "total, actionable, critical, warning, byType"},
    ]
    print(json.dumps({"week": f"{mon}..{sun}", "missingWeekMeta": missing,
                      "note": "Run deck_week.py plan/save/build first; it pulls this week's totals, which also fill the 4-week table. "
                              "If missingWeekMeta is not empty, pull get_account_metrics groupBy none for each of those Mondays..Sundays "
                              "and save with deck_week.py save meta <monday> <file>.",
                      "write": f"{b.tmp}/pulled.json with one key per call below, then deck.py extra-build {slug} {b.tmp}/pulled.json",
                      "calls": calls}, indent=1))


def week_row(meta, mon, last):
    r = {"label": label(mon), "revenue": meta["revenue"], "units": meta["units"], "orders": meta["orders"],
         "impressions": meta["impr"], "clicks": meta["clicks"], "refundUnits": meta.get("refundUnits"),
         "refundRate": meta["refundRate"], "spend": meta["spend"], "adSales": meta["adSales"], "acos": meta["acos"], "tacos": meta["tacos"]}
    if last:
        r.update(net=meta["net"], refunds=meta["refunds"], adOrders=meta["adOrders"], cpc=meta["cpc"],
                 cvr=round(meta["adOrders"] / meta["clicks"] * 100, 2) if meta["clicks"] else 0)
    return r


def cmd_extra_build(slug, pulled_path):
    b = Brand(slug); M = mondays(b); mon = M[-1]
    wkdir = os.path.join(b.deck, "wk")
    P = json.load(open(pulled_path))
    miss = [str(m) for m in M if not os.path.exists(os.path.join(wkdir, f"meta_{m}.json"))]
    if miss: die("weekly totals missing (see extra-plan)", weeks=miss)
    metas = [json.load(open(os.path.join(wkdir, f"meta_{m}.json"))) for m in M]
    if any(x.get("refundUnits") is None for x in metas): die("a saved week has no refundUnits; re-save it with deck_week.py save meta")
    weeks = [week_row(x, m, i == 3) for i, (x, m) in enumerate(zip(metas, M))]
    daily = P["daily"]
    if len(daily) == 7:
        prior = json.load(open(os.path.join(b.deck, "weeks", str(mon - dt.timedelta(days=1)), "extra.json")))["daily"][7:]
        daily = prior + daily
    if len(daily) != 14: die(f"need 14 daily rows, have {len(daily)}")
    want = [str(mon + dt.timedelta(days=i - 7)) for i in range(14)]
    if [r[0] for r in daily] != want: die("daily rows are not the two Mon–Sun weeks", have=[r[0] for r in daily])
    for k in ("mtd", "mtdLY", "searchTerms", "targets", "reviews", "alerts"):
        if k not in P: die(f"pulled.json is missing {k}")
    bj = json.load(open(os.path.join(b.deck, "brand.json"))) if os.path.exists(os.path.join(b.deck, "brand.json")) else {}
    X = {"_source": f"Sellerise MCP, account {bj.get('account', 'Dormeo')}, marketplace {bj.get('marketplace', 'US')}, "
                    f"pulled {dt.date.today()} (Monday-Sunday weeks; 4-week table from the saved weekly totals)",
         "weeks": weeks, "mtd": P["mtd"], "mtdLY": P["mtdLY"], "daily": daily,
         "searchTerms": P["searchTerms"], "targets": P["targets"], "reviews": P["reviews"], "alerts": P["alerts"]}
    out = os.path.join(b.tmp, "extra.json"); json.dump(X, open(out, "w"), indent=1, ensure_ascii=False)
    print(json.dumps({"status": "ok", "extra": out, "weeks": [w["label"] for w in weeks], "revenue": weeks[-1]["revenue"]}))


def cmd_build(slug):
    b = Brand(slug); ex, nt = os.path.join(b.tmp, "extra.json"), os.path.join(b.tmp, "notes.json")
    ref = os.path.join(b.weeks()[-1], "notes.json") if b.weeks() else os.path.join(HERE, "example", "notes.json")
    gone = sorted(set(json.load(open(ref))) - set(json.load(open(nt))))
    if gone: die("notes.json is missing keys", keys=gone)
    dd, px = build_deck(b, ex, nt)
    from PIL import Image                                            # one contact sheet instead of 17 image reads
    ims = [Image.open(f) for f in sorted(glob.glob(os.path.join(b.tmp, "s-*.jpg")))]
    w, h = 480, int(480 * ims[0].height / ims[0].width); cols = 3; rows = -(-len(ims) // cols)
    sheet = Image.new("RGB", (cols * w + (cols + 1) * 8, rows * h + (rows + 1) * 8), "white")
    for i, im in enumerate(ims):
        sheet.paste(im.resize((w, h)), (8 + (i % cols) * (w + 8), 8 + (i // cols) * (h + 8)))
    cs = os.path.join(b.tmp, "contact.jpg"); sheet.save(cs, quality=80)
    print(json.dumps({"status": "built", "pptx": px, "contactSheet": cs, "slides": len(ims),
                      "tip": "look at contactSheet; open a single s-NN.jpg only if something looks off"}))


def cmd_publish(slug):
    b = Brand(slug); dd = os.path.join(b.tmp, "deck_data.json"); W = json.load(open(dd))["week"]
    b.run([sys.executable, "publish_deck.py", os.path.join(b.tmp, "deck.pptx"), dd, b.site])
    wd = os.path.join(b.deck, "weeks", W["end"]); os.makedirs(wd, exist_ok=True)
    for f in ("extra.json", "notes.json"): shutil.copy(os.path.join(b.tmp, f), wd)
    c = commit_push(b, f"{b.e['brand']} weekly deck: week {W['num']} ({W['label']})")
    print(json.dumps({"status": "published", "brand": b.e["brand"], "week": W["num"], "commit": c, "googleSlidesId": b.e.get("googleSlidesId"),
                      "next": "nothing: the Apps Script rebuilds the editable Google Slides file within ~30 minutes"}))


def cmd_slides(slug, ids, only):
    b = Brand(slug)
    a = [sys.executable, "gslides_images.py", os.path.join(b.site, "deck"), ids] + (["--only", only] if only else [])
    sys.stdout.write(b.run(a))


if __name__ == "__main__":
    a = sys.argv[1:]
    if not a: print(__doc__); sys.exit(1)
    c = a[0]
    if c == "brands": cmd_brands()
    elif c == "index": cmd_index()
    elif c == "tuesday": cmd_tuesday(a[1], "--dry-run" in a)
    elif c == "extra-plan": cmd_extra_plan(a[1])
    elif c == "extra-build": cmd_extra_build(a[1], a[2])
    elif c == "build": cmd_build(a[1])
    elif c == "publish": cmd_publish(a[1])
    elif c == "slides": cmd_slides(a[1], a[2], a[a.index("--only") + 1] if "--only" in a else None)
    else: print(__doc__); sys.exit(1)
